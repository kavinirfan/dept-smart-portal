from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.seed import DEMO_PASSWORD


def auth(client, email):
    r = client.post("/api/auth/login", data={"username": email, "password": DEMO_PASSWORD})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def user_id(client, email):
    return client.get("/api/auth/me", headers=auth(client, email)).json()["user_id"]


def subject_id(client, code):
    subs = client.get("/api/subjects", headers=auth(client, "hod@dept.edu")).json()
    return next(s["subject_id"] for s in subs if s["subject_code"] == code)


# ---- Marks ------------------------------------------------------------------

def test_teacher_sees_only_credentialed_subjects(client):
    subs = client.get("/api/marks/my-subjects", headers=auth(client, "arjun@dept.edu")).json()
    assert [s["subject_code"] for s in subs] == ["CS501", "CS504"]


def test_teacher_without_credential_is_blocked(client):
    h = auth(client, "lakshmi@dept.edu")  # holds CS502 only
    r = client.post("/api/marks/components", headers=h, json={
        "subject_id": subject_id(client, "CS501"), "name": "Quiz", "question_max": [5], "weightage": 5})
    assert r.status_code == 403


def test_enter_marks_and_student_sees_breakdown(client):
    h = auth(client, "arjun@dept.edu")
    comp = client.post("/api/marks/components", headers=h, json={
        "subject_id": subject_id(client, "CS501"), "name": "Quiz 1", "question_max": [5, 5], "weightage": 5}).json()
    kavin = user_id(client, "kavin@dept.edu")
    bad = client.put(f"/api/marks/sheet/{comp['component_id']}", headers=h,
                     json=[{"student_id": kavin, "question_marks": [6, 1]}])
    assert bad.status_code == 422  # above the question maximum
    ok = client.put(f"/api/marks/sheet/{comp['component_id']}", headers=h,
                    json=[{"student_id": kavin, "question_marks": [4, 3.5]}])
    assert ok.json() == {"saved": 1}

    report = client.get(f"/api/marks/student/{kavin}", headers=auth(client, "kavin@dept.edu")).json()
    dbms = next(s for s in report["subjects"] if s["subject"]["subject_code"] == "CS501")
    quiz = next(c for c in dbms["components"] if c["name"] == "Quiz 1")
    assert quiz["question_marks"] == [4, 3.5] and quiz["marks_obtained"] == 7.5 and quiz["weighted"] == 3.75


def test_student_cannot_see_other_students(client):
    priya = user_id(client, "priya@dept.edu")
    assert client.get(f"/api/marks/student/{priya}", headers=auth(client, "kavin@dept.edu")).status_code == 403


def test_marks_export(client):
    r = client.get(f"/api/export/marks?subject_id={subject_id(client, 'CS501')}", headers=auth(client, "arjun@dept.edu"))
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames[0] == "Internal summary" and "CAT-1" in wb.sheetnames


# ---- Labs -------------------------------------------------------------------

def _slot(days=2, start=10, hours=1):
    s = datetime.combine(date.today() + timedelta(days=days), datetime.min.time()) + timedelta(hours=start)
    return s.isoformat(), (s + timedelta(hours=hours)).isoformat()


def test_first_come_first_served(client):
    labs = client.get("/api/labs", headers=auth(client, "arjun@dept.edu")).json()
    lab = labs[0]["lab_id"]
    start, end = _slot()
    first = client.post("/api/labs/bookings", headers=auth(client, "arjun@dept.edu"),
                        json={"lab_id": lab, "start_time": start, "end_time": end, "purpose": "Extra session"})
    assert first.status_code == 201
    # Overlapping request from someone else loses.
    s2, e2 = _slot(start=10, hours=2)
    second = client.post("/api/labs/bookings", headers=auth(client, "lakshmi@dept.edu"),
                         json={"lab_id": lab, "start_time": s2, "end_time": e2, "purpose": "Clash"})
    assert second.status_code == 409 and "Arjun" in second.json()["detail"]
    # Back-to-back is fine.
    s3, e3 = _slot(start=11)
    assert client.post("/api/labs/bookings", headers=auth(client, "lakshmi@dept.edu"),
                       json={"lab_id": lab, "start_time": s3, "end_time": e3, "purpose": "Next slot"}).status_code == 201
    # After cancelling, the slot is free again.
    bid = first.json()["booking_id"]
    assert client.delete(f"/api/labs/bookings/{bid}", headers=auth(client, "lakshmi@dept.edu")).status_code == 403
    assert client.delete(f"/api/labs/bookings/{bid}", headers=auth(client, "arjun@dept.edu")).status_code == 204
    assert client.post("/api/labs/bookings", headers=auth(client, "lakshmi@dept.edu"),
                       json={"lab_id": lab, "start_time": start, "end_time": end, "purpose": "Now free"}).status_code == 201


def test_students_cannot_book_and_past_rejected(client):
    lab = client.get("/api/labs", headers=auth(client, "arjun@dept.edu")).json()[0]["lab_id"]
    start, end = _slot(days=3)
    assert client.post("/api/labs/bookings", headers=auth(client, "kavin@dept.edu"),
                       json={"lab_id": lab, "start_time": start, "end_time": end, "purpose": "x"}).status_code == 403
    ps, pe = _slot(days=-2)
    assert client.post("/api/labs/bookings", headers=auth(client, "arjun@dept.edu"),
                       json={"lab_id": lab, "start_time": ps, "end_time": pe, "purpose": "x"}).status_code == 422


def test_maintenance_blocks_booking_and_only_incharge_sets_status(client):
    lab = client.get("/api/labs", headers=auth(client, "arjun@dept.edu")).json()[2]["lab_id"]
    assert client.patch(f"/api/labs/{lab}/status", headers=auth(client, "arjun@dept.edu"), json={"status": "maintenance"}).status_code == 403
    assert client.patch(f"/api/labs/{lab}/status", headers=auth(client, "labincharge@dept.edu"), json={"status": "maintenance"}).status_code == 200
    start, end = _slot(days=4)
    r = client.post("/api/labs/bookings", headers=auth(client, "arjun@dept.edu"),
                    json={"lab_id": lab, "start_time": start, "end_time": end, "purpose": "x"})
    assert r.status_code == 409
    client.patch(f"/api/labs/{lab}/status", headers=auth(client, "labincharge@dept.edu"), json={"status": "free"})


# ---- Faculty ------------------------------------------------------------------

def test_faculty_status_update_and_board(client):
    h = auth(client, "rahul@dept.edu")
    assert client.post("/api/faculty/status", headers=h, json={"status": "Available", "note": "Room 101"}).status_code == 200
    board = client.get("/api/faculty", headers=auth(client, "kavin@dept.edu")).json()
    rahul = next(f for f in board if f["name"] == "Prof. Rahul Menon")
    assert rahul["status"] == "Available" and rahul["note"] == "Room 101"
    assert client.post("/api/faculty/status", headers=auth(client, "kavin@dept.edu"), json={"status": "Available"}).status_code == 403


def test_faculty_export_admin_only(client):
    assert client.get("/api/export/faculty", headers=auth(client, "arjun@dept.edu")).status_code == 403
    r = client.get("/api/export/faculty", headers=auth(client, "hod@dept.edu"))
    assert load_workbook(BytesIO(r.content)).sheetnames == ["Attendance", "Status log"]


# ---- Notices ------------------------------------------------------------------

def test_notices_post_filter_delete(client):
    h = auth(client, "kavin@dept.edu")
    before = client.get("/api/notices", headers=h).json()
    assert all(not n["expired"] for n in before)  # expired hidden by default
    n = client.post("/api/notices", headers=h, json={
        "title": "Hackathon volunteers", "company": "IEEE", "description": "Paid internship", "type": "internship",
        "deadline": (date.today() + timedelta(days=5)).isoformat(), "link": "https://ieee.org"}).json()
    found = client.get("/api/notices?q=hackathon&type=internship", headers=h).json()
    assert [x["notice_id"] for x in found] == [n["notice_id"]]
    assert client.delete(f"/api/notices/{n['notice_id']}", headers=auth(client, "priya@dept.edu")).status_code == 403
    assert client.delete(f"/api/notices/{n['notice_id']}", headers=h).status_code == 204
