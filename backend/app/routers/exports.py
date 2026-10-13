"""Excel (.xlsx) exports. The database stays authoritative; each download is generated fresh from it."""
from datetime import date, datetime, time, timedelta
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import FacultyState, FacultyStatus, LabBooking, Mark, MarkComponent, Notice, Role, Subject, User
from ..security import get_current_user, require_roles
from .marks import require_credential

router = APIRouter(prefix="/api/export", tags=["exports"])
HEADER_FILL = PatternFill("solid", fgColor="DCE4F7")


def add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    ws.freeze_panes = "A2"
    return ws


def send(wb: Workbook, filename: str) -> StreamingResponse:
    if len(wb.sheetnames) > 1:
        del wb[wb.sheetnames[0]]  # the default empty sheet
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/marks")
def export_marks(subject_id: int, db: Session = Depends(get_db), user: User = Depends(require_roles(Role.teacher, Role.admin))):
    subject = db.get(Subject, subject_id)
    if subject is None:
        raise HTTPException(status_code=404, detail="Subject not found")
    require_credential(db, user, subject_id)
    comps = db.scalars(select(MarkComponent).where(MarkComponent.subject_id == subject_id).order_by(MarkComponent.component_id)).all()
    students = db.scalars(select(User).where(User.role == Role.student).order_by(User.roll_no)).all()
    marks = {(m.student_id, m.component_id): m for m in db.scalars(
        select(Mark).where(Mark.component_id.in_([c.component_id for c in comps])))}

    wb = Workbook()
    summary = []
    for s in students:
        row, internal = [s.roll_no, s.name], 0.0
        for c in comps:
            m = marks.get((s.user_id, c.component_id))
            if m is None:
                row.append(None)
            elif m.absent:
                row.append("AB")
            else:
                w = m.marks_obtained / c.max_marks * c.weightage
                internal += w
                row.append(round(w, 2))
        summary.append(row + [round(internal, 2)])
    add_sheet(wb, "Internal summary",
              ["Roll no", "Name"] + [f"{c.name} (/{c.weightage:g})" for c in comps] + [f"Internal (/{sum(c.weightage for c in comps):g})"],
              summary)
    for c in comps:
        rows = []
        for s in students:
            m = marks.get((s.user_id, c.component_id))
            qs = m.question_marks if m else [None] * len(c.question_max)
            rows.append([s.roll_no, s.name, *qs, ("AB" if m.absent else m.marks_obtained) if m else None])
        add_sheet(wb, c.name, ["Roll no", "Name"] + [f"Q{i} (/{q:g})" for i, q in enumerate(c.question_max, 1)] + [f"Total (/{c.max_marks:g})"], rows)
    return send(wb, f"marks_{subject.subject_code}.xlsx")


@router.get("/labs")
def export_labs(days: int = 30, db: Session = Depends(get_db), _: User = Depends(require_roles(Role.teacher, Role.lab_incharge, Role.admin))):
    since = datetime.combine(date.today() - timedelta(days=days), time.min)
    rows = db.scalars(select(LabBooking).where(LabBooking.start_time >= since).order_by(LabBooking.start_time)).all()
    wb = Workbook()
    add_sheet(wb, "Lab bookings", ["Booking ID", "Lab", "Date", "Start", "End", "Booked by", "Purpose", "Status", "Requested at (UTC)"], [
        [b.booking_id, b.lab.lab_name, b.start_time.date().isoformat(), b.start_time.strftime("%H:%M"), b.end_time.strftime("%H:%M"),
         b.booker.name, b.purpose, "Cancelled" if b.cancelled else "Active", b.created_at.strftime("%Y-%m-%d %H:%M:%S")]
        for b in rows
    ])
    return send(wb, f"lab_bookings_{date.today().isoformat()}.xlsx")


@router.get("/faculty")
def export_faculty(days: int = 30, db: Session = Depends(get_db), _: User = Depends(require_roles(Role.admin))):
    since = datetime.combine(date.today() - timedelta(days=days), time.min)
    rows = db.scalars(select(FacultyStatus).where(FacultyStatus.updated_at >= since).order_by(FacultyStatus.updated_at)).all()
    wb = Workbook()
    # Attendance-style grid: a faculty member counts as present on a day if they reported Available or In Class.
    present_states = {FacultyState.available, FacultyState.in_class}
    days_list = sorted({r.updated_at.date() for r in rows})
    faculty = sorted({r.faculty for r in rows}, key=lambda f: f.name)
    grid = []
    for f in faculty:
        mine = [r for r in rows if r.faculty_id == f.user_id]
        line = [f.name]
        for d in days_list:
            states = [r.status for r in mine if r.updated_at.date() == d]
            line.append("P" if any(s in present_states for s in states) else ("L" if FacultyState.on_leave in states else ("-" if not states else "A")))
        grid.append(line + [sum(1 for v in line[1:] if v == "P")])
    add_sheet(wb, "Attendance", ["Faculty"] + [d.strftime("%d %b") for d in days_list] + ["Days present"], grid)
    add_sheet(wb, "Status log", ["When (UTC)", "Faculty", "Status", "Note"], [
        [r.updated_at.strftime("%Y-%m-%d %H:%M"), r.faculty.name, r.status.value, r.note] for r in rows
    ])
    return send(wb, f"faculty_attendance_{date.today().isoformat()}.xlsx")


@router.get("/notices")
def export_notices(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    rows = db.scalars(select(Notice).order_by(Notice.posted_on.desc())).all()
    wb = Workbook()
    add_sheet(wb, "Notices", ["Type", "Title", "Company", "Deadline", "Link", "Posted by", "Posted on"], [
        [n.type.value, n.title, n.company, n.deadline.isoformat() if n.deadline else None, n.link, n.poster.name,
         n.posted_on.strftime("%Y-%m-%d")] for n in rows
    ])
    return send(wb, "notices.xlsx")
